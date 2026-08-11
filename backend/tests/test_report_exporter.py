"""Unit tests for app.utils.report_exporter's Defects sheet - specifically
that Confidence and Verification Status are surfaced as their own columns
instead of being dropped at export time or conflated into Severity."""
from openpyxl import load_workbook
import io

from app.utils.report_exporter import build_excel_report

_BASE_DEFECT = {
    "defect_id": "DEF-1-001", "test_case_id": "TC-1-001", "plan_step_id": 1,
    "page_title": "Home", "page_url": "https://example.com/",
    "principle": "Perceivable", "criteria": "1.1.1 Non-text Content", "level": "A",
    "element_type": "Images / Icons / Image Buttons",
    "status": "fail", "defect_status": "Open",
    "description": "desc", "expected_result": "expected", "actual_result": "actual",
    "steps_to_reproduce": "steps", "suggestion_to_fix": "fix it", "html_snippet": "<img>",
    "screenshot": None,
}


def _defect(**overrides):
    d = dict(_BASE_DEFECT)
    d.update(overrides)
    return d


def test_defects_sheet_has_verification_status_and_confidence_columns():
    defects = [
        _defect(severity="Critical", status="fail", verification_status="Confirmed", confidence=1.0),
        _defect(defect_id="DEF-1-002", severity="Review Needed", status="needs_review",
                verification_status="AI Judgment - Needs Review", confidence=0.4),
    ]
    xlsx_bytes = build_excel_report({"test_cases": [], "defects": defects})
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Defects"]

    headers = [c.value for c in ws[1]]
    assert "Verification Status" in headers
    assert "Confidence" in headers
    assert "Severity" in headers
    # Confidence and Verification Status must be separate columns from Severity,
    # not folded into it - this is the fix for confirmed/needs_review conflation.
    assert headers.index("Confidence") != headers.index("Severity")
    assert headers.index("Verification Status") != headers.index("Severity")

    vs_col = headers.index("Verification Status") + 1
    conf_col = headers.index("Confidence") + 1
    row2 = [ws.cell(row=2, column=c).value for c in (vs_col, conf_col)]
    row3 = [ws.cell(row=3, column=c).value for c in (vs_col, conf_col)]
    assert row2 == ["Confirmed", "100%"]
    assert row3 == ["AI Judgment - Needs Review", "40%"]


def test_defects_sheet_review_needed_severity_has_a_color():
    """'Review Needed' used to have no _SEVERITY_BG entry, so it silently
    blended into the default row background - regression guard."""
    from app.utils.report_exporter import _SEVERITY_BG
    assert "Review Needed" in _SEVERITY_BG
