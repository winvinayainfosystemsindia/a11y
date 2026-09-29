"""
Builds the IAAP-style "Accessibility Conformance Certificate" Word report by
filling the firm's own template (app/utils/templates/accessibility_
conformance_certificate_template.docx) with a specific audit run's real
data, rather than generating a generic report from scratch.

The template file is itself a previously-completed example report (title
page, evaluation methodology, an "Applicable Standards" table, a "Terms"
glossary, and two large per-success-criterion tables - Level A and Level AA -
each row stating Supports/Partially Supports/Does Not Support/Not Applicable/
Not Evaluated plus a remarks paragraph). This module edits the cover-page
fields in place, rewrites the methodology/environment text to accurately
describe *this platform's* pipeline (not the template's original manual-
tester tools, which would misrepresent what was actually done), corrects the
"Applicable Standards" table to only claim what this platform actually
evaluates, and replaces the two sample criteria tables with freshly built
ones populated from the run's real test-case results - while reusing the
template's exact fonts/borders/shading so the output still looks like the
firm's own report, not a generic docx.
"""
from __future__ import annotations

import copy
import io
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from app.utils.wcag_criteria import WCAG_LEVEL_A, WCAG_LEVEL_AA_ADDITIONAL, level_for_criterion

_TEMPLATE_PATH = Path(__file__).parent / "templates" / "accessibility_conformance_certificate_template.docx"

_HEADER_SHADING = "AEAAAA"
_BORDER_COLOR = "000000"
_CELL_FONT_SIZE = 12

_METHODS_TEXT = (
    "The accessibility evaluation was carried out by an AI-driven automated audit agent that combines "
    "deterministic rule-based testing with AI-assisted semantic judgment for comprehensive coverage. "
    "Deterministic testing used axe-core (an open-source, industry-standard WCAG rule-conformance engine) for "
    "structural/markup checks, a custom WCAG-formula contrast-ratio calculator for color contrast, and an "
    "automated keyboard tab-order and focus-trap simulator for keyboard accessibility. For criteria requiring "
    "contextual or semantic understanding that cannot be verified by rule-based testing alone (e.g. whether "
    "alternative text is meaningful, or an error message is understandable), an AI judgment model evaluated the "
    "criterion directly against the page content, with a confidence score attached to every finding; any finding "
    "below the confidence threshold is flagged in this report as requiring manual review rather than asserted as "
    "conclusive."
)

_ENV_LINES = [
    "Testing Environment: Automated, server-side evaluation via a headless browser - not a live manual "
    "assistive-technology session.",
    "Browser Engine: Chromium (via Playwright 1.41.2)",
    "Automated Rule Engine: axe-core v4.12.1 (deterministic WCAG rule checks)",
    "Custom Deterministic Checks: WCAG contrast-ratio calculation; keyboard tab-order and focus-trap simulation",
    "AI-Assisted Judgment: large-language-model semantic analysis for criteria requiring contextual "
    "understanding, with confidence scoring and automatic flagging for manual review when confidence is low",
]

# Standard/Guideline row label (matched by prefix against the table's own
# first-column text, not a numeric row index - robust to the template
# reordering its rows) -> corrected "Included In Report" cell text. This
# platform evaluates WCAG 2.0/2.1/2.2 Level A+AA (see app/utils/
# wcag_criteria.py) - it must not claim the India-specific
# BIS/GIGWA standards, which it does not evaluate.
_STANDARDS_CORRECTIONS = {
    "BIS IS17802": "Not Evaluated",
    "GIGWA": "Not Evaluated",
}


def _find_para(doc: Document, predicate) -> Paragraph | None:
    for p in doc.paragraphs:
        if predicate(p):
            return p
    return None


def _para_index(doc: Document, target: Paragraph) -> int:
    """doc.paragraphs rebuilds a fresh list of Paragraph wrappers on every
    access, and Paragraph doesn't define __eq__, so list.index(target) fails
    with a false "not in list" whenever target came from an earlier access -
    compare by the underlying XML element instead."""
    for i, p in enumerate(doc.paragraphs):
        if p._p is target._p:
            return i
    raise ValueError("paragraph not found in document")


def _set_para_text(para: Paragraph, text: str) -> None:
    """Replace a paragraph's visible text while keeping its first run's
    formatting (font/size/bold/style) and dropping any extra runs, so the
    template's exact look survives even though the content is now dynamic."""
    if not para.runs:
        para.add_run(text)
        return
    para.runs[0].text = text
    for extra in para.runs[1:]:
        extra.text = ""


def _clone_para_after(para: Paragraph) -> Paragraph:
    new_p = copy.deepcopy(para._p)
    para._p.addnext(new_p)
    return Paragraph(new_p, para._parent)


def _delete_para(para: Paragraph) -> None:
    para._p.getparent().remove(para._p)


def _elements_between(body, start_el, end_el) -> list:
    collecting = False
    found = []
    for el in body:
        if el is start_el:
            collecting = True
            continue
        if end_el is not None and el is end_el:
            break
        if collecting:
            found.append(el)
    return found


# ── Cell/table styling (reverse-engineered from the template's own tables so
# freshly-built replacement tables match its borders/shading/font exactly) ──

def _apply_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), _BORDER_COLOR)
        borders.append(el)
    tblPr.append(borders)


def _shade_cell(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def _fill_cell(cell, text: str, *, bold: bool = False, shading: str | None = None, center: bool = True) -> None:
    # cell.text = "" leaves one empty run behind (it clears content, then adds
    # a paragraph with a single empty run) - add_run() would just append a
    # second run after it rather than replacing it, so the formatting below
    # would land on the wrong (second) run. Reuse that first run directly.
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = para.runs[0]
    run.text = text
    run.bold = bold
    run.font.size = Pt(_CELL_FONT_SIZE)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if shading:
        _shade_cell(cell, shading)


# ── Per-criterion conformance rollup ────────────────────────────────────────

def _rollup_conformance(rows_for_criterion: list[dict]) -> str:
    """One WCAG criterion is usually tested via several element-level test
    cases (per-instance enumeration) on a real page - roll their individual
    pass/fail/needs_review verdicts up into one of the template's five
    defined conformance terms."""
    statuses = {r["status"] for r in rows_for_criterion}
    if not statuses:
        return "Not Evaluated"
    has_fail = "fail" in statuses
    has_pass = "pass" in statuses
    has_review = "needs_review" in statuses
    if has_fail and has_pass:
        return "Partially Supports"
    if has_fail:
        return "Does Not Support"
    if has_review:
        return "Not Evaluated"
    return "Supports"


def _short(text: str | None, limit: int = 170) -> str:
    text = " ".join((text or "").split())
    if len(text) <= limit:
        return text
    return text[:limit].rsplit(" ", 1)[0] + "…"


def _remarks_for_criterion(name: str, rows: list[dict], conformance: str, *, in_scope: bool) -> str:
    if not rows:
        if not in_scope:
            return f"This criterion is outside the conformance level targeted by this audit and was not evaluated."
        return "This criterion was not covered by any test case in this audit."

    if conformance == "Supports":
        n = len(rows)
        return (
            f"All {n} evaluated instance{'s' if n != 1 else ''} of this criterion on this page met the "
            f"requirement."
        )

    if conformance in ("Does Not Support", "Partially Supports"):
        failing = [r for r in rows if r["status"] == "fail"]
        passing = [r for r in rows if r["status"] == "pass"]
        parts = []
        if conformance == "Partially Supports":
            parts.append(f"{len(passing)} of {len(rows)} evaluated instance(s) met this requirement; "
                         f"{len(failing)} did not.")
        else:
            parts.append(f"{len(failing)} of {len(rows)} evaluated instance(s) did not meet this requirement.")
        for r in failing[:2]:
            detail = _short(r.get("remediation") or r.get("actual_result"))
            parts.append(f"'{r['tc_name']}': {detail}")
        return " ".join(parts)

    # Not Evaluated
    reviews = [r for r in rows if r["status"] == "needs_review"]
    n = len(reviews)
    detail = _short(reviews[0].get("description")) if reviews else ""
    base = (
        f"{n} instance{'s' if n != 1 else ''} of this criterion could not be conclusively evaluated by "
        f"automated/AI analysis and require manual verification."
    )
    return f"{base} {detail}" if detail else base


def _build_criteria_table(doc: Document, criteria_list: list[tuple[str, str]], rows_by_criterion: dict,
                           conformance_level: str) -> "Table":  # noqa: F821 - docx.table.Table, avoid extra import
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    _apply_table_borders(table)

    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, ["Criteria", "Conformance Level", "Remarks and Explanations"]):
        _fill_cell(cell, text, bold=True, shading=_HEADER_SHADING)
    _set_repeat_header(table.rows[0])

    for sc_id, name in criteria_list:
        rows_for = rows_by_criterion.get(sc_id, [])
        sc_level = level_for_criterion(sc_id)
        in_scope = sc_level == "A" or conformance_level != "A"
        conformance = _rollup_conformance(rows_for) if (rows_for or in_scope) else "Not Evaluated"
        remarks = _remarks_for_criterion(name, rows_for, conformance, in_scope=in_scope)

        cells = table.add_row().cells
        _fill_cell(cells[0], f"{sc_id} {name} (Level {sc_level})")
        _fill_cell(cells[1], conformance)
        _fill_cell(cells[2], remarks, center=False)

    for row in table.rows:
        row.cells[0].width = Inches(4.68)
        row.cells[1].width = Inches(2.03)
        row.cells[2].width = Inches(3.44)

    return table


def _replace_criteria_section(doc: Document, heading_text_prefix: str, next_heading_text_prefix: str | None,
                               criteria_list: list[tuple[str, str]], rows_by_criterion: dict,
                               conformance_level: str) -> None:
    heading = _find_para(doc, lambda p: p.text.strip().startswith(heading_text_prefix))
    if heading is None:
        return
    next_heading = (
        _find_para(doc, lambda p: p.text.strip().startswith(next_heading_text_prefix))
        if next_heading_text_prefix else None
    )

    idx = _para_index(doc, heading)
    # The template always has "Notes: ..." then a blank spacer paragraph
    # immediately after each section heading - keep both, replace everything
    # from just after the spacer up to (not including) the next heading.
    notes_para = doc.paragraphs[idx + 1]
    spacer_para = doc.paragraphs[idx + 2]

    end_el = next_heading._p if next_heading is not None else None
    for el in _elements_between(doc.element.body, spacer_para._p, end_el):
        el.getparent().remove(el)

    table = _build_criteria_table(doc, criteria_list, rows_by_criterion, conformance_level)
    spacer_para._p.addnext(table._tbl)


def _fix_standards_table(doc: Document) -> None:
    table = doc.tables[0]
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        for prefix, correction in _STANDARDS_CORRECTIONS.items():
            if label.startswith(prefix):
                _fill_cell(row.cells[1], correction, center=False)
                break


def build_conformance_certificate(results: dict, *, client_name: str | None = None,
                                   report_date: str | None = None) -> bytes:
    """The Word report the "Download IAAP Word Report" button produces -
    the firm's own Accessibility Conformance Certificate template, filled
    with this audit run's real per-criterion results."""
    doc = Document(str(_TEMPLATE_PATH))

    client = client_name or results.get("project_name") or results.get("page_title") or "Client"
    page_title = results.get("page_title") or results.get("page_url", "")
    conformance_level = results.get("conformance_level", "AA")
    date_str = report_date or datetime.now().strftime("%B %d, %Y")

    # ── Cover page ───────────────────────────────────────────────────────
    if doc.paragraphs:
        _set_para_text(doc.paragraphs[0], client)  # "Name of the Firm" -> client/site name

    product_para = _find_para(doc, lambda p: p.text.strip().startswith("Name of Product"))
    if product_para is not None:
        _set_para_text(product_para, f"Name of Product: {page_title}")

    date_para = _find_para(doc, lambda p: p.text.strip().startswith("Report Date"))
    if date_para is not None:
        _set_para_text(date_para, f"Report Date: {date_str}")

    client_para = _find_para(doc, lambda p: p.text.strip().rstrip(":") == "Client")
    if client_para is not None:
        _set_para_text(client_para, f"Client: {client}")

    # ── Evaluation Methods Used ─────────────────────────────────────────
    methods_heading = _find_para(doc, lambda p: p.text.strip() == "Evaluation Methods Used:")
    if methods_heading is not None:
        idx = _para_index(doc, methods_heading)
        _set_para_text(doc.paragraphs[idx + 1], _METHODS_TEXT)

    # ── Evaluation Environment ──────────────────────────────────────────
    env_bullets = [
        p for p in doc.paragraphs
        if p.style.name == "List Paragraph"
        and p.text.strip().startswith(("Operating System", "Browser", "Screen Reader"))
    ]
    if env_bullets:
        last = env_bullets[-1]
        while len(env_bullets) < len(_ENV_LINES):
            last = _clone_para_after(last)
            env_bullets.append(last)
        for p, text in zip(env_bullets, _ENV_LINES):
            _set_para_text(p, text)
        for extra in env_bullets[len(_ENV_LINES):]:
            _delete_para(extra)

    # The template's static "Page 1 of 18" line is wrong for any report of a
    # different length than the original example - remove it rather than
    # print a misleading page count.
    page_num_para = _find_para(doc, lambda p: re.match(r"^Page \d+ of \d+$", p.text.strip()))
    if page_num_para is not None:
        _delete_para(page_num_para)

    # ── Applicable Standards/Guidelines ─────────────────────────────────
    _fix_standards_table(doc)

    # ── Table 1 (Level A) / Table 2 (Level AA) ──────────────────────────
    rows_by_criterion: dict[str, list[dict]] = defaultdict(list)
    for row in results.get("test_cases", []):
        sc_id = (row.get("criteria") or "").split(" ", 1)[0]
        if sc_id:
            rows_by_criterion[sc_id].append(row)

    _replace_criteria_section(doc, "Table 1: Success Criteria, Level A", "Table 2: Success Criteria, Level AA",
                               WCAG_LEVEL_A, rows_by_criterion, conformance_level)
    _replace_criteria_section(doc, "Table 2: Success Criteria, Level AA", None,
                               WCAG_LEVEL_AA_ADDITIONAL, rows_by_criterion, conformance_level)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
